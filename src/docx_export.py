"""Deterministic, read-only DOCX rendering for Odysseus documents."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone

from docx import Document


_EMAIL_ENTRY_MARKER = "<!-- odysseus-email-entry -->"
_FIXED_TIMESTAMP = datetime(2000, 1, 1, tzinfo=timezone.utc)
_ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)


def _add_content(document: Document, content: str) -> None:
    """Render a small, predictable Markdown subset while preserving line order."""
    for raw_line in (content or "").replace(_EMAIL_ENTRY_MARKER, "").splitlines():
        line = raw_line.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        unordered = re.match(r"^\s*[-*+]\s+(.+)$", line)
        ordered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        elif unordered:
            document.add_paragraph(unordered.group(1), style="List Bullet")
        elif ordered:
            document.add_paragraph(ordered.group(1), style="List Number")
        else:
            document.add_paragraph(line)


def _normalize_package(package: bytes) -> bytes:
    """Normalize ZIP metadata so equivalent input produces equivalent bytes."""
    source = io.BytesIO(package)
    output = io.BytesIO()
    with zipfile.ZipFile(source, "r") as source_zip, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as output_zip:
        for name in source_zip.namelist():
            original = source_zip.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=_ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            output_zip.writestr(info, source_zip.read(name))
    return output.getvalue()


def render_document_docx(title: str, content: str) -> bytes:
    """Return a DOCX representation without mutating persistent state."""
    document = Document()
    document.core_properties.created = _FIXED_TIMESTAMP
    document.core_properties.modified = _FIXED_TIMESTAMP
    document.core_properties.title = title or "Untitled"
    document.add_heading(title or "Untitled", level=0)
    _add_content(document, content)

    buffer = io.BytesIO()
    document.save(buffer)
    return _normalize_package(buffer.getvalue())

import io
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import core.database as cdb
import routes.document_routes as document_routes
import src.database as src_database
from core.database import Document, DocumentVersion
from core.database import Session as DbSession
from src.agent_tools.document_tools import CreateDocumentTool, set_active_document
from src.docx_export import render_document_docx


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MARKER = "<!-- odysseus-email-entry -->"


@pytest.fixture
def document_db(tmp_path, monkeypatch):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'documents.db'}",
        connect_args={"check_same_thread": False},
    )
    cdb.Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    monkeypatch.setattr(document_routes, "SessionLocal", session_factory)
    monkeypatch.setattr(src_database, "SessionLocal", session_factory)
    yield session_factory
    engine.dispose()


def _request(user):
    return SimpleNamespace(state=SimpleNamespace(current_user=user))


def _endpoint(method, path):
    router = document_routes.setup_document_routes(SimpleNamespace(), None)
    for route in router.routes:
        if route.path == path and method in route.methods:
            return route.endpoint
    raise AssertionError(f"Missing endpoint: {method} {path}")


def _seed_document(session_factory):
    content = "\n".join(
        [
            "## Ordered messages",
            MARKER,
            "### 2026-01-01 — First synthetic email",
            "First body",
            MARKER,
            "### 2026-01-02 — Second synthetic email",
            "Second body",
        ]
    )
    db = session_factory()
    try:
        db.add(DbSession(id="alice-session", owner="alice", name="Alice", model="m", endpoint_url="http://x"))
        db.add(Document(
            id="alice-doc",
            session_id="alice-session",
            title='Unsafe / Quarterly: "Mail"',
            language="markdown",
            current_content=content,
            version_count=1,
            is_active=True,
            owner="alice",
        ))
        db.add(DocumentVersion(
            id="alice-version",
            document_id="alice-doc",
            version_number=1,
            content=content,
            summary="Initial version",
            source="user",
        ))
        db.commit()
    finally:
        db.close()
    return content


def test_renderer_is_valid_reopenable_and_deterministic():
    content = f"# Heading\n{MARKER}\n- Alpha\n1. Beta"

    first = render_document_docx("Export title", content)
    second = render_document_docx("Export title", content)

    assert first == second
    with zipfile.ZipFile(io.BytesIO(first)) as package:
        assert "word/document.xml" in package.namelist()
        assert MARKER.encode() not in package.read("word/document.xml")
    reopened = WordDocument(io.BytesIO(first))
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    assert "Export title" in text
    assert "Heading" in text
    assert "Alpha" in text
    assert "Beta" in text
    assert MARKER not in text


@pytest.mark.asyncio
async def test_export_endpoint_returns_current_docx_without_mutation(document_db):
    original_content = _seed_document(document_db)
    export = _endpoint("GET", "/api/document/{doc_id}/export-docx")

    response = await export(_request("alice"), "alice-doc")

    assert response.status_code == 200
    assert response.media_type == DOCX_MIME
    disposition = response.headers["content-disposition"]
    assert disposition == 'attachment; filename="Unsafe_Quarterly_Mail.docx"'
    assert disposition.endswith('.docx"')
    with zipfile.ZipFile(io.BytesIO(response.body)) as package:
        document_xml = package.read("word/document.xml")
        assert MARKER.encode() not in document_xml

    reopened = WordDocument(io.BytesIO(response.body))
    visible_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    first = "2026-01-01 — First synthetic email"
    second = "2026-01-02 — Second synthetic email"
    assert "Unsafe / Quarterly" in visible_text
    assert visible_text.index(first) < visible_text.index(second)
    assert visible_text.count(first) == 1
    assert visible_text.count(second) == 1
    assert MARKER not in visible_text

    db = document_db()
    try:
        stored = db.query(Document).filter(Document.id == "alice-doc").one()
        versions = db.query(DocumentVersion).filter(DocumentVersion.document_id == "alice-doc").all()
        assert stored.current_content == original_content
        assert stored.version_count == 1
        assert len(versions) == 1
        assert versions[0].content == original_content
    finally:
        db.close()


@pytest.mark.asyncio
async def test_export_endpoint_preserves_owner_scope(document_db):
    _seed_document(document_db)
    export = _endpoint("GET", "/api/document/{doc_id}/export-docx")

    with pytest.raises(HTTPException) as exc:
        await export(_request("bob"), "alice-doc")

    assert exc.value.status_code == 404


@pytest.mark.asyncio
async def test_create_document_tool_exposes_docx_url(document_db):
    db = document_db()
    try:
        db.add(DbSession(id="tool-session", owner="alice", name="Tool", model="m", endpoint_url="http://x"))
        db.commit()
    finally:
        db.close()

    try:
        result = await CreateDocumentTool().execute(
            "Tool title\nmarkdown\nTool body",
            {"session_id": "tool-session", "owner": "alice"},
        )
    finally:
        set_active_document(None)

    assert result["action"] == "create"
    assert result["title"] == "Tool title"
    assert result["content"] == "Tool body"
    assert result["version"] == 1
    assert result["docx_url"] == f"/api/document/{result['doc_id']}/export-docx"
